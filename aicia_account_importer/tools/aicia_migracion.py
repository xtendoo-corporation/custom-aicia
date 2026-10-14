#!/usr/bin/env python3
# Copyright 2026 Xtendoo - https://www.xtendoo.es/
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).
"""Utilidades de migración AICIAWIN → Odoo, repetibles para cada ejercicio.

Tres comandos (ver README.md):

  terceros       General.mdb → Proveedores.xlsx y Clientes.xlsx (los lee el importador
                 de terceros de Odoo).
  asiento-socio  Facturas del año + Apuntes/Lineas → asiento_socio.xlsx (fichero
                 opcional del importador de apuntes: socio exacto de cada asiento).
  conciliacion   AICIAWIN contra lo importado en Odoo → informe de conciliación.

Solo usa la librería estándar y openpyxl. Para leer el .mdb necesita mdbtools
(``mdb-export``); también se puede pasar una carpeta con las tablas en CSV.
"""

import argparse
import collections
import csv
import datetime
import io
import re
import subprocess
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    openpyxl = None

csv.field_size_limit(10**9)

# Cuentas legadas que, por criterio contable, no llevan socio (ver el módulo).
NO_PARTNER = {
    "430000000", "401000000", "400000000", "436000000", "616000000", "610000000",
    "465000000", "401000200", "401000300", "401000400", "618000200",
    "460000273", "460001012", "460001118",
}
# Cuentas especiales que se importan con los 9 dígitos.
FULL_CODE = {"401000200", "401000300", "401000400", "618000200"}
# Grupo de cuenta → tipo de tercero (S proveedor, C cliente, E personal).
GROUP = {
    "400": "S", "401": "S", "410": "S", "430": "C", "431": "C", "436": "C",
    "460": "E", "465": "E", "610": "E", "611": "E", "616": "E", "618": "E",
}
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ── Lectura de datos ─────────────────────────────────────────────────────────

def read_mdb(mdb, table):
    """Filas de una tabla del .mdb (requiere mdb-export)."""
    out = subprocess.run(
        ["mdb-export", "-D", "%Y-%m-%d", mdb, table],
        check=True, capture_output=True, text=True, encoding="utf-8", errors="replace",
    ).stdout
    return list(csv.DictReader(io.StringIO(out)))


def read_source(args, table):
    if getattr(args, "csv_dir", None):
        with open(f"{args.csv_dir}/{table}.csv", encoding="utf-8", errors="replace") as fh:
            return list(csv.DictReader(fh))
    return read_mdb(args.mdb, table)


def read_xlsx(path):
    """Primera hoja de un xlsx como lista de dicts (todo como texto)."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    head = [str(c) for c in next(rows)]
    return [
        {h: ("" if v is None else (str(v) if not isinstance(v, bool) else str(v))) for h, v in zip(head, row)}
        for row in rows
    ]


def read_csv(path):
    with open(path, encoding="utf-8", errors="replace") as fh:
        return list(csv.DictReader(fh))


def is_true(value):
    return str(value).strip().lower() in ("true", "1", "-1")


def date_ymd(value):
    """'20250113' o datetime/'2025-01-13 00:00:00' → 'YYYYMMDD'."""
    text = re.sub(r"\D", "", str(value))[:8]
    return text


def to_int(value):
    try:
        return int(round(float(str(value).replace(",", "."))))
    except ValueError:
        return 0


# ── Comando: terceros ────────────────────────────────────────────────────────

INT_COLUMNS = {"ID_Proveedor", "ID_Cliente", "ID_Pais", "Fecha_Alta", "ID_PaisEnv"}
BOOL_COLUMNS = {"ProfesionalLibre", "IRPFReducido", "IRPFProfEspecial", "Embargo", "Socio_Protector"}


def cmd_terceros(args):
    for table in ("Proveedores", "Clientes"):
        rows = read_source(args, table)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = table
        head = list(rows[0].keys())
        ws.append(head)
        for row in rows:
            out = []
            for col in head:
                value = row[col]
                if value == "":
                    out.append(None)
                elif col in INT_COLUMNS:
                    out.append(to_int(value))
                elif col in BOOL_COLUMNS:
                    out.append(value in ("1", "True", "true", "-1"))
                else:
                    out.append(ILLEGAL.sub("", value))
            ws.append(out)
        path = f"{args.out}/{table}.xlsx"
        wb.save(path)
        print(f"{table}: {len(rows)} filas → {path}")


# ── Comando: asiento-socio ───────────────────────────────────────────────────

def build_asiento_socio(apuntes, facturas_prov, facturas_cli, year):
    """Devuelve (filas, estadísticas). Cada fila: (Numero_Apunte, Tipo, ID_Tercero)."""
    by_id = {a["ID_Apunte"]: a for a in apuntes}
    stats = collections.Counter()
    result = {}
    for tipo, facturas, key in (
        ("proveedor", facturas_prov, "ID_Proveedor"),
        ("cliente", facturas_cli, "ID_Cliente"),
    ):
        for f in facturas:
            if str(f.get("Año_Apunte")) != str(year):
                continue
            stats[f"facturas {tipo}"] += 1
            a = by_id.get(f["ID_Apunte"])
            if not a:
                stats[f"factura {tipo} sin apunte"] += 1
                continue
            stats["fecha coincide"] += date_ymd(f["Fecha_Factura"]) == date_ymd(a["Fecha_Contable"])
            stats["importe coincide"] += (
                to_int(f["Base"]) + to_int(f["Cuota"]) == to_int(a["Importe_Total"])
            )
            k = (to_int(a["Numero_Apunte"]), tipo)
            if k in result and result[k] != to_int(f[key]):
                stats["asiento con dos terceros del mismo tipo"] += 1
            result.setdefault(k, to_int(f[key]))
    rows = [(n, t, i) for (n, t), i in sorted(result.items())]
    return rows, stats


def cmd_asiento_socio(args):
    apuntes = read_xlsx(args.apuntes)
    rows, stats = build_asiento_socio(
        apuntes, read_source(args, "Fact_Proveedores"), read_source(args, "Fact_Clientes"), args.year
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Socio por asiento"
    ws.append(["Numero_Apunte", "Tipo", "ID_Tercero", "Origen"])
    for n, t, i in rows:
        ws.append([n, t, i, "factura"])
    wb.save(args.out)
    print(f"{len(rows)} asientos con socio por factura → {args.out}")
    for k, v in sorted(stats.items()):
        print(f"  {k}: {v}")
    total = stats["facturas proveedor"] + stats["facturas cliente"]
    if total and stats["fecha coincide"] < 0.95 * total:
        print("AVISO: la fecha de la factura coincide con la del apunte en menos del 95 %: "
              "revisa que Apuntes y Fact_* son del mismo ejercicio.", file=sys.stderr)


# ── Comando: conciliacion ────────────────────────────────────────────────────

def target_account(legacy):
    return legacy if legacy in FULL_CODE else legacy[:5] + "0"


def expected_lines(apuntes, lineas, facturas, owners):
    """Líneas esperadas por asiento: {num: [(cuenta Odoo, debe, haber, socio, motivo, cuenta legada, fecha)]}.

    ``socio`` es None o (tipo, número). ``facturas`` es {(num, tipo S/C): id}.
    ``owners`` es {tipo: {cuenta legada: [ids con esa cuenta]}}.
    """
    valid = {a["ID_Apunte"]: a for a in apuntes if is_true(a["Validado"]) and not is_true(a["Anulado"])}
    result = collections.defaultdict(list)
    for l in lineas:
        a = valid.get(l["ID_Apunte"])
        if not a:
            continue
        legacy = l["Cuenta_Contable"]
        num = str(to_int(a["Numero_Apunte"]))
        amount = to_int(l["Importe"])
        debit, credit = (amount, 0) if l["Tipo_Contable"] == "D" else (0, amount)
        kind = GROUP.get(legacy[:3])
        who, why = None, "sin tipo de socio"
        if kind:
            suffix = int(legacy[-4:])
            if kind in "SC" and (num, kind) in facturas:
                who, why = (kind, facturas[(num, kind)]), "por factura"
            elif legacy in NO_PARTNER:
                who, why = None, "sin socio por regla contable"
            elif kind == "E":
                who, why = ("E", suffix), "por código de personal"
            else:
                members = owners[kind].get(legacy, [])
                if len(members) == 1:
                    who, why = (kind, members[0]), "cuenta de un solo socio"
                elif len(members) > 1:
                    who, why = (kind, suffix), "cuenta compartida: titular"
                else:
                    who, why = (kind, suffix), "cuenta antigua sin socio hoy: código = ID"
        result[num].append((target_account(legacy), debit, credit, who, why, legacy, date_ymd(a["Fecha_Contable"])))
    return result


def partner_matches(expected, got):
    """¿El socio esperado es el que tiene la línea de Odoo? Un contacto puede tener varios códigos."""
    if expected is None:
        return not (got["cli"] or got["prov"] or got["emp"])
    kind, number = expected
    return str(number) in {"C": got["cli"], "S": got["prov"], "E": got["emp"]}[kind]


def reconcile(expected, odoo_rows):
    """Empareja líneas esperadas y de Odoo por asiento. Devuelve (ok, difs, solo_esperadas, solo_odoo)."""
    got = collections.defaultdict(list)
    for r in odoo_rows:
        codes = {k: set(filter(None, (r.get(k) or "").split("|"))) for k in ("cli", "prov", "emp")}
        got[r["num"]].append(
            {"key": (r["acc"], int(r["debit"]), int(r["credit"])), "used": False, **codes}
        )
    matched, wrong_partner, missing = [], [], []
    for num, lines in expected.items():
        pool = got.get(num, [])
        for line in lines:
            key = line[:3]
            cands = [g for g in pool if not g["used"] and g["key"] == key]
            if not cands:
                missing.append((num, line))
                continue
            best = next((g for g in cands if partner_matches(line[3], g)), None)
            if best:
                best["used"] = True
                matched.append((num, line))
            else:
                cands[0]["used"] = True
                wrong_partner.append((num, line))
    extra = [(num, g["key"]) for num, pool in got.items() for g in pool if not g["used"]]
    return matched, wrong_partner, missing, extra


def cmd_conciliacion(args):
    apuntes = read_xlsx(args.apuntes)
    lineas = read_xlsx(args.lineas)
    odoo = read_csv(args.odoo_lines)
    asiento_socio = {}
    if args.asiento_socio:
        for r in read_xlsx(args.asiento_socio):
            asiento_socio[(str(to_int(r["Numero_Apunte"])), "S" if r["Tipo"].startswith("prov") else "C")] = to_int(r["ID_Tercero"])
    owners = {"S": collections.defaultdict(list), "C": collections.defaultdict(list)}
    for r in read_source(args, "Proveedores"):
        owners["S"][r["Cuenta_Contable"]].append(int(r["ID_Proveedor"]))
    for r in read_source(args, "Clientes"):
        owners["C"][r["Cuenta_Contable"]].append(int(r["ID_Cliente"]))
    expected = expected_lines(apuntes, lineas, asiento_socio, owners)
    matched, wrong, missing, extra = reconcile(expected, odoo)
    write_conciliation(args.out, expected, odoo, matched, wrong, missing, extra)


def write_conciliation(path, expected, odoo, matched, wrong, missing, extra):
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    grey = PatternFill("solid", fgColor="D9D9D9")
    wb = openpyxl.Workbook()

    def sheet(title, head, rows):
        ws = wb.create_sheet(title)
        ws.append(head)
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows():
            for c in row:
                c.font = font
        for c in ws[1]:
            c.font, c.fill = bold, grey
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(50, max(10, max(len(str(c.value or "")) for c in col[:80]) + 2))
        ws.freeze_panes = "A2"
        return ws

    # Cuentas: AICIAWIN contra Odoo
    sums_a, sums_o = collections.defaultdict(lambda: [0, 0, 0]), collections.defaultdict(lambda: [0, 0, 0])
    for lines in expected.values():
        for acc, d, h, *_ in lines:
            s = sums_a[acc]; s[0] += d; s[1] += h; s[2] += 1
    for r in odoo:
        s = sums_o[r["acc"]]; s[0] += int(r["debit"]); s[1] += int(r["credit"]); s[2] += 1
    accounts = sorted(set(sums_a) | set(sums_o))
    ws_c = sheet(
        "Cuentas",
        ["Cuenta Odoo", "Debe AICIAWIN (€)", "Haber AICIAWIN (€)", "Líneas AICIAWIN", "Debe Odoo (€)", "Haber Odoo (€)", "Líneas Odoo", "Dif. Debe", "Dif. Haber", "Dif. líneas"],
        [[a, sums_a[a][0] / 100, sums_a[a][1] / 100, sums_a[a][2], sums_o[a][0] / 100, sums_o[a][1] / 100, sums_o[a][2],
          f"=B{i}-E{i}", f"=C{i}-F{i}", f"=D{i}-G{i}"] for i, a in enumerate(accounts, start=2)],
    )
    last = len(accounts) + 1
    ws_c.append(["TOTAL"] + [f"=SUM({c}2:{c}{last})" for c in "BCDEFGHIJ"])
    for c in ws_c[last + 1]:
        c.font = bold
    # Socio por motivo
    reasons = collections.Counter(line[4] for _n, line in matched)
    certainty = {
        "por factura": "Dato", "por código de personal": "Dato (código de la cuenta)",
        "cuenta de un solo socio": "Dato", "sin socio por regla contable": "Decisión del contable",
        "cuenta compartida: titular": "Inferido", "cuenta antigua sin socio hoy: código = ID": "Inferido",
        "sin tipo de socio": "No aplica (la cuenta no lleva socio)",
    }
    ws_s = sheet("Socio por motivo", ["Motivo", "Certeza", "Líneas"],
                 [[k, certainty.get(k, ""), v] for k, v in sorted(reasons.items(), key=lambda x: -x[1])])
    ws_s.append(["Líneas con socio distinto del esperado", "Pendiente", len(wrong)])
    ws_s.append(["TOTAL", "", f"=SUM(C2:C{ws_s.max_row})"])
    for c in ws_s[ws_s.max_row]:
        c.font = bold
    # Detalles
    def fecha(f): return f"{f[6:8]}/{f[4:6]}/{f[:4]}" if len(f) == 8 else f
    sheet("Líneas inferidas", ["Nº asiento", "Fecha", "Cuenta legada", "Debe (€)", "Haber (€)", "Socio asignado", "Motivo"],
          [[int(n), fecha(l[6]), l[5], l[1] / 100, l[2] / 100, f"{l[3][0]}{l[3][1]}" if l[3] else "", l[4]]
           for n, l in sorted(matched, key=lambda x: (int(x[0]), x[1][5])) if l[4].startswith(("cuenta compartida", "cuenta antigua"))])
    sheet("Pendientes y diferencias", ["Nº asiento", "Fecha", "Cuenta legada", "Debe (€)", "Haber (€)", "Socio esperado", "Qué pasa"],
          [[int(n), fecha(l[6]), l[5], l[1] / 100, l[2] / 100, f"{l[3][0]}{l[3][1]}" if l[3] else "", "El socio de Odoo no es el esperado (o no tiene socio)"] for n, l in wrong]
          + [[int(n), fecha(l[6]), l[5], l[1] / 100, l[2] / 100, "", "Falta la línea en Odoo"] for n, l in missing]
          + [[int(n), "", "", "", "", "", f"Línea de más en Odoo {k}"] for n, k in extra])
    # Resumen
    ws = wb.active
    ws.title = "Resumen"
    n_exp = sum(len(v) for v in expected.values())
    rows = [
        ["Conciliación AICIAWIN → Odoo", None, None],
        ["Generado", datetime.date.today().strftime("%d/%m/%Y"), None],
        [],
        ["", "AICIAWIN", "Odoo", "Diferencia"],
        ["Asientos", len(expected), len({r["num"] for r in odoo}), "=B5-C5"],
        ["Líneas", n_exp, len(odoo), "=B6-C6"],
        ["Debe (€)", "=Cuentas!B%d" % (last + 1), "=Cuentas!E%d" % (last + 1), "=B7-C7"],
        ["Haber (€)", "=Cuentas!C%d" % (last + 1), "=Cuentas!F%d" % (last + 1), "=B8-C8"],
        ["Cuentas con diferencia de saldo", None, None,
         f'=SUMPRODUCT((ABS(Cuentas!H2:H{last})+ABS(Cuentas!I2:I{last})>0.001)*1)'],
        [],
        ["Líneas con el socio esperado", len(matched), None, None],
        ["…de ellas inferidas (cuenta compartida o cuenta antigua)",
         sum(1 for _n, l in matched if l[4].startswith(("cuenta compartida", "cuenta antigua"))), None, None],
        ["Líneas con socio distinto o sin socio esperado", len(wrong), None, None],
        ["Líneas que faltan en Odoo", len(missing), None, None],
        ["Líneas de más en Odoo", len(extra), None, None],
        [],
        ["Criterio: un asiento está conciliado si sus líneas (cuenta de Odoo, debe y haber) son idénticas a las de AICIAWIN. "
         "Las 'inferidas' son líneas sin factura en cuentas compartidas o en cuentas que ningún socio usa hoy; "
         "no se pueden confirmar con los datos de origen."],
    ]
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.font = font
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    for c in ws[4]:
        c.font, c.fill = bold, grey
    ws.column_dimensions["A"].width = 62
    for col in "BCD":
        ws.column_dimensions[col].width = 16
    wb.save(path)
    print(f"{len(matched)} líneas conciliadas, {len(wrong)} con socio distinto, {len(missing)} faltan, {len(extra)} de más → {path}")


# ── CLI ──────────────────────────────────────────────────────────────────────

def main(argv=None):
    if openpyxl is None:
        sys.exit("Falta openpyxl (pip install openpyxl).")
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    def source(sp):
        g = sp.add_mutually_exclusive_group(required=True)
        g.add_argument("--mdb", help="General.mdb de AICIAWIN (requiere mdbtools)")
        g.add_argument("--csv-dir", help="carpeta con las tablas en CSV (mdb-export)")

    s = sub.add_parser("terceros", help="General.mdb → Proveedores.xlsx y Clientes.xlsx")
    source(s)
    s.add_argument("--out", required=True, help="carpeta de salida")
    s.set_defaults(func=cmd_terceros)

    s = sub.add_parser("asiento-socio", help="fichero asiento → socio a partir de las facturas")
    source(s)
    s.add_argument("--apuntes", required=True, help="Apuntes.xlsx del ejercicio")
    s.add_argument("--year", required=True, type=int, help="ejercicio (Año_Apunte), p. ej. 2025")
    s.add_argument("--out", required=True, help="asiento_socio.xlsx")
    s.set_defaults(func=cmd_asiento_socio)

    s = sub.add_parser("conciliacion", help="informe de conciliación AICIAWIN ↔ Odoo")
    source(s)
    s.add_argument("--apuntes", required=True)
    s.add_argument("--lineas", required=True, help="Lineas_Apunte.xlsx del ejercicio")
    s.add_argument("--odoo-lines", required=True, help="CSV de odoo_lines.sql")
    s.add_argument("--asiento-socio", help="el mismo asiento_socio.xlsx que se importó")
    s.add_argument("--out", required=True, help="informe .xlsx")
    s.set_defaults(func=cmd_conciliacion)

    args = p.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
